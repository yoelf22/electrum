#!/usr/bin/env python3
"""Export one run's documents as a single Word file.

    python3 scripts/build_docx.py output/<slug> [out.docx]

Cover page, then a table of contents, then every phase document in order with the
PNGs embedded. The contents is a Word TOC field: Word fills it on open (say Yes to
"update fields"); Google Docs and Pages show its placeholder line instead.
Needs python-docx.
"""
import datetime
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from mdblocks import parse, runs  # noqa: E402

try:
    from docx import Document
    from docx.enum.text import WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
except ImportError:
    sys.exit("python-docx is missing: pip install python-docx")

from build_html import ORDER, product_name  # noqa: E402

GREY = RGBColor(0x6B, 0x74, 0x79)


def toc_field(doc):
    """Word TOC field over Heading 1–2, filled by Word when the file is opened."""
    p = doc.add_paragraph()
    r = p.add_run()
    for kind, text in (("begin", None), ("instr", 'TOC \\o "1-2" \\h \\z \\u'), ("separate", None),
                       ("text", "Open in Word and choose Yes when asked to update fields; "
                                "this line becomes the table of contents."), ("end", None)):
        if kind == "instr":
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = text
        elif kind == "text":
            el = OxmlElement("w:t"); el.text = text
        else:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), kind)
        r._r.append(el)
    upd = OxmlElement("w:updateFields"); upd.set(qn("w:val"), "true")
    doc.settings.element.append(upd)


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def render(doc, blocks, folder):
    for kind, level, payload in blocks:
        if kind == "h":
            if level == 1:
                continue  # the document's own title is the Heading 1 we already wrote
            doc.add_heading(re.sub(r"\*\*|`", "", payload), level=min(level, 4))
        elif kind == "p":
            runs(doc.add_paragraph(), payload)
        elif kind == "prompt":
            r = runs(doc.add_paragraph(), payload)
            for x in r.runs: x.italic, x.font.color.rgb = True, GREY
        elif kind == "bullet":
            runs(doc.add_paragraph(style="List Bullet 2" if level else "List Bullet"), payload)
        elif kind == "check":
            runs(doc.add_paragraph(style="List Bullet"), payload)
        elif kind == "num":
            runs(doc.add_paragraph(style="List Number"), re.sub(r"^\d+\.\s+", "", payload))
        elif kind == "code":
            lang, body = payload
            p = doc.add_paragraph(); r = p.add_run("\n".join(body)); r.font.name, r.font.size = "Consolas", Pt(8.5)
        elif kind == "table":
            rows = payload
            if not rows: continue
            w = max(len(r) for r in rows)
            tbl = doc.add_table(rows=len(rows), cols=w); tbl.style = "Table Grid"
            for i, row in enumerate(rows):
                for j in range(w):
                    cell = tbl.cell(i, j); cell.text = ""
                    runs(cell.paragraphs[0], row[j] if j < len(row) else "", bold=(i == 0))
                    for x in cell.paragraphs[0].runs: x.font.size = Pt(8.5)
            doc.add_paragraph()
        elif kind == "img":
            p = folder / payload
            if p.exists():
                doc.add_picture(str(p), width=Cm(16))
            else:
                runs(doc.add_paragraph(), f"[image: {payload}]")


def build(folder, out):
    files = [folder / n for n in ORDER if (folder / n).exists()]
    if not files:
        sys.exit(f"no phase files in {folder}")
    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2.2)
    labels = [f.stem.replace("_", " ").title() for f in files]
    # cover
    doc.add_paragraph().paragraph_format.space_before = Pt(120)
    r = doc.add_paragraph().add_run("ELECTRUM"); r.bold = True; r.font.color.rgb = RGBColor(0x8A, 0x6D, 0x1F)
    doc.add_paragraph(product_name(folder), style="Title")
    r = doc.add_paragraph().add_run("Product definition"); r.font.size, r.font.color.rgb = Pt(14), GREY
    doc.add_paragraph(datetime.date.today().isoformat())
    doc.add_paragraph()
    doc.add_paragraph("Documents in this file: " + ", ".join(labels) + ".")
    page_break(doc)
    doc.add_heading("Contents", level=1)
    toc_field(doc)
    for f, label in zip(files, labels):
        page_break(doc)
        doc.add_heading(label, level=1)
        render(doc, parse(f.read_text(encoding="utf-8")), folder)
    doc.save(out)
    print(f"{out} ({len(files)} documents, {out.stat().st_size // 1024} KB)")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    folder = Path(sys.argv[1])
    out = Path(sys.argv[2]) if len(sys.argv) > 2 else folder / f"{folder.name}.docx"
    build(folder, out)
