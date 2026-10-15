"""Build fill-in Word forms from the templates, pre-filled with a worked example.

Template headings and guiding questions stay as fixed text. Every answer area is a
Word content control whose grey placeholder is the example's content; it disappears
when the user types or pastes into it. The checklist gets tickable boxes and a
PASS/FAIL/N/A dropdown.

    python scripts/build_word_forms.py OUT_DIR [EXAMPLE_DIR]

Needs python-docx and Pillow. Only Word honours content controls; Google Docs and Pages show
the placeholders as plain grey text.
"""
import difflib
import io
import itertools
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

ROOT = Path(__file__).resolve().parent.parent
TPL = ROOT / "templates"
INK, GUIDE, ACCENT = RGBColor(0x1B, 0x23, 0x28), RGBColor(0x6B, 0x74, 0x79), RGBColor(0x8A, 0x6D, 0x1F)
PLACEHOLDER = RGBColor(0x9A, 0xA0, 0xA6)
_ids = itertools.count(1000)

# ---------------------------------------------------------------- markdown

def parse(text):
    """Markdown -> list of blocks: (kind, level, payload)."""
    lines, out, i = text.split("\n"), [], 0
    while i < len(lines):
        l = lines[i]
        if l.startswith("```"):
            lang, body = l[3:].strip(), []
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                body.append(lines[i]); i += 1
            out.append(("code", 0, (lang, body)))
        elif l.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                if not re.match(r"^\|[\s:|-]+\|?\s*$", lines[i]):
                    rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            out.append(("table", 0, rows)); continue
        elif m := re.match(r"^(#{1,4})\s+(.*)$", l):
            out.append(("h", len(m[1]), m[2].strip()))
        elif m := re.match(r"^>\s?-?\s*(.*)$", l):
            if m[1].strip(): out.append(("prompt", 0, m[1]))
        elif m := re.match(r"^!\[[^\]]*\]\(([^)]+)\)", l):
            out.append(("img", 0, m[1]))
        elif m := re.match(r"^(\s*)- \[[ xX]\]\s+(.*)$", l):
            out.append(("check", 0, m[2]))
        elif m := re.match(r"^(\s*)[-*]\s+(.*)$", l):
            out.append(("bullet", len(m[1]) // 2, m[2]))
        elif m := re.match(r"^\s*(\d+)\.\s+(.*)$", l):
            out.append(("num", 0, f"{m[1]}. {m[2]}"))
        elif re.match(r"^---+\s*$", l):
            pass
        elif m := re.match(r"^\*([^*].*)\*$", l.strip()):
            out.append(("prompt", 0, m[1]))
        elif l.strip():
            out.append(("p", 0, l.strip()))
        i += 1
    return out


def sections(blocks):
    """Split blocks into [heading, level, body] lists; body = blocks before the next heading."""
    secs = [[None, 0, []]]
    for b in blocks:
        if b[0] == "h": secs.append([b[2], b[1], []])
        else: secs[-1][2].append(b)
    return secs


def norm(h):
    h = re.sub(r"\[[^\]]*\]|\([^)]*\)|\*\*|[—–:]", " ", h.lower())
    return " ".join(h.split())


def key(h):
    n = norm(h)
    if m := re.match(r"^(\d+(\.\d+)?)\.?\s", n): return m[1]
    if m := re.match(r"^(scenario \d+|decision \d+)", n): return m[1]
    return n


def match(tsecs, esecs):
    """Map template heading index -> example section index (same level, closest name)."""
    out, used = {}, set()
    for ti, (th, tl, _) in enumerate(tsecs):
        if th is None: continue
        best, score = None, 0.0
        for ei, (eh, el, _) in enumerate(esecs):
            if eh is None or el != tl or ei in used: continue
            a, b = key(th), key(eh)
            s = 1.0 if a == b else 0.0 if re.match(r"^[\d.]+$|^(scenario|decision) \d", a + " " + b) and a[:3] != b[:3] \
                else difflib.SequenceMatcher(None, a, b).ratio()
            if a and b and (a in b or b in a): s = max(s, 0.8)
            if s > score: best, score = ei, s
        if best is not None and score >= 0.55:
            out[ti] = best; used.add(best)
    return out

# ---------------------------------------------------------------- docx helpers

def runs(par, text, placeholder=False, **base):
    """Add **bold**, *italic*, `code` runs to a paragraph."""
    for tok in re.split(r"(\*\*[^*]+\*\*|`[^`]+`|(?<![\w*])\*[^*\s][^*]*\*(?!\w)|(?<!\w)_[^_]+_(?!\w))", text):
        if not tok: continue
        fmt = dict(base)
        if tok.startswith("**"): tok, fmt["bold"] = tok[2:-2], True
        elif tok.startswith("`"): tok, fmt["font"] = tok[1:-1], "Consolas"
        elif tok[0] in "*_" and tok[-1] == tok[0] and len(tok) > 2: tok, fmt["italic"] = tok[1:-1], True
        r = par.add_run(tok)
        if placeholder: r.style = "Placeholder Text"
        if fmt.get("bold"): r.bold = True
        if fmt.get("italic"): r.italic = True
        if fmt.get("font"): r.font.name = fmt["font"]
        if fmt.get("color"): r.font.color.rgb = fmt["color"]
        if fmt.get("size"): r.font.size = fmt["size"]
    return par


def sdt(alias, inline=False, dropdown=None, checkbox=False):
    """Return (sdt, sdtContent). Rich text showing placeholder unless dropdown/checkbox."""
    el, pr, content = OxmlElement("w:sdt"), OxmlElement("w:sdtPr"), OxmlElement("w:sdtContent")
    a = OxmlElement("w:alias"); a.set(qn("w:val"), alias[:60]); pr.append(a)
    i = OxmlElement("w:id"); i.set(qn("w:val"), str(next(_ids))); pr.append(i)
    if not checkbox: pr.append(OxmlElement("w:showingPlcHdr"))
    if dropdown:
        dd = OxmlElement("w:dropDownList")
        for v in dropdown:
            li = OxmlElement("w:listItem"); li.set(qn("w:displayText"), v); li.set(qn("w:value"), v); dd.append(li)
        pr.append(dd)
    if checkbox:
        cb = OxmlElement("w14:checkbox")
        c = OxmlElement("w14:checked"); c.set(qn("w14:val"), "0"); cb.append(c)
        for tag, code in (("w14:checkedState", "2612"), ("w14:uncheckedState", "2610")):
            s = OxmlElement(tag); s.set(qn("w14:val"), code); s.set(qn("w14:font"), "MS Gothic"); cb.append(s)
        pr.append(cb)
    el.append(pr); el.append(content)
    return el, content


def faded(path):
    """Wash an example image out toward white so it reads as a placeholder."""
    img = Image.open(path).convert("RGB")
    buf = io.BytesIO()
    Image.blend(img, Image.new("RGB", img.size, "white"), 0.65).save(buf, "PNG")
    buf.seek(0)
    return buf


class Form:
    def __init__(self):
        self.doc = d = Document()
        sec = d.sections[0]
        sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Cm(2.1)
        st = d.styles
        st["Normal"].font.name, st["Normal"].font.size, st["Normal"].font.color.rgb = "Calibri", Pt(11), INK
        st["Normal"].paragraph_format.space_after = Pt(5)
        for name, size, color in (("Title", 24, INK), ("Heading 1", 17, INK), ("Heading 2", 14, INK),
                                  ("Heading 3", 12, ACCENT), ("Heading 4", 11, INK)):
            s = st[name]; s.font.name, s.font.size, s.font.bold, s.font.color.rgb = "Cambria", Pt(size), True, color
            s.font.italic = False
            if name == "Title": s.paragraph_format.space_after = Pt(12)
        ph = st.add_style("Placeholder Text", WD_STYLE_TYPE.CHARACTER)
        ph.font.color.rgb = PLACEHOLDER
        g = st.add_style("Guidance", WD_STYLE_TYPE.PARAGRAPH)
        g.base_style, g.font.italic, g.font.size, g.font.color.rgb = st["Normal"], True, Pt(9.5), GUIDE
        g.paragraph_format.space_after = Pt(2); g.paragraph_format.left_indent = Cm(0.5)
        self.target = None  # sdtContent to divert body paragraphs into
        self.new_page = False  # next heading starts a page
        self.about = None  # "About this document" paragraph to place under the next heading

    def _place(self, el):
        if self.target is not None: self.target.append(el)

    def para(self, text="", style=None, placeholder=False, **fmt):
        p = self.doc.add_paragraph(style=style)
        runs(p, text, placeholder, **fmt)
        self._place(p._p)
        return p

    def heading(self, text, level, product=None):
        p = self.doc.add_heading(level=level)
        if self.new_page:  # an empty spacer left before the break would strand a blank page
            prev = p._p.getprevious()
            if prev is not None and prev.tag == qn("w:p") and not "".join(prev.itertext()).strip():
                prev.getparent().remove(prev)
        p.paragraph_format.page_break_before, self.new_page = self.new_page, False
        if product and "[Product Name]" in text:
            pre, post = text.split("[Product Name]", 1)
            runs(p, pre); self.inline(p, product, "Product name"); runs(p, post)
        else:
            runs(p, text)
        if self.about:
            self.para(f"**About this document.** {self.about}"); self.about = None
        return p

    def inline(self, par, text, alias, dropdown=None):
        el, content = sdt(alias, inline=True, dropdown=dropdown)
        tmp = self.doc.add_paragraph(); runs(tmp, text, placeholder=True)
        for r in list(tmp._p.findall(qn("w:r"))): content.append(r)
        tmp._p.getparent().remove(tmp._p)
        par._p.append(el)

    def block(self, blocks, alias, base):
        """One block-level content control holding example blocks as placeholder."""
        el, content = sdt(alias)
        self.doc.element.body.insert(len(self.doc.element.body) - 1, el)
        self.target = content
        for kind, lvl, val in blocks:
            if kind == "img":
                p = self.para(); p.add_run().add_picture(faded(base / val), width=Cm(14))
            elif kind == "code":
                lang, body = val
                if lang == "mermaid": self.para("[Diagram: sketch here or paste an image]", placeholder=True)
                else:
                    for b in body: self.para(b or " ", placeholder=True, font="Consolas", size=Pt(8.5)).paragraph_format.space_after = Pt(0)
            elif kind == "bullet":
                self.para(val, style="List Bullet" + (f" {lvl + 1}" if lvl else ""), placeholder=True)
            elif kind == "h":
                self.para(f"**{val}**", placeholder=True)
            else:
                self.para(val, placeholder=True)
        self.target = None

    def table(self, rows, alias, fixed_cols=0):
        """Header row fixed; body cells are inline placeholders (first `fixed_cols` columns fixed)."""
        cols = max(len(r) for r in rows)
        t = self.doc.add_table(rows=len(rows), cols=cols)
        t.style = "Table Grid"
        for ri, row in enumerate(rows):
            for ci in range(cols):
                cell, txt = t.cell(ri, ci), (row[ci] if ci < len(row) else "")
                p = cell.paragraphs[0]
                if ri == 0 or ci < fixed_cols: runs(p, txt, bold=ri == 0)
                elif txt: self.inline(p, txt, f"{alias} · {rows[0][ci] if ci < len(rows[0]) else ''}")
                else: self.inline(p, " ", alias)
        self.doc.add_paragraph()
        return t

# ---------------------------------------------------------------- body rendering

GUIDE_KINDS = ("prompt",)


def is_guidance(b):
    """Template prose with no fill-in brackets reads as instructions, not an answer."""
    return b[0] in GUIDE_KINDS or (b[0] == "p" and "[" not in b[2] and not b[2].startswith("**"))


def render_body(f, tbody, ebody, title, base):
    """Guidance from the template, then the answer area from the example (or the template's own fill-ins)."""
    for b in tbody:
        if is_guidance(b): f.para(b[2], style="Guidance")
    answer = ebody if ebody is not None else [b for b in tbody if not is_guidance(b)]
    chunk = []

    def flush():
        if chunk: f.block(chunk, title, base); chunk.clear()

    for b in answer:
        kind, _, val = b
        label = re.match(r"^\*\*([^*]{1,60}?:)\*\*\s*(.*)$", val) if kind == "p" else None
        if kind == "table":
            flush(); f.table(val, title)
        elif label:
            flush()
            p = f.para(f"**{label[1]}**")
            if label[2]: p.add_run(" "); f.inline(p, label[2], label[1].rstrip(":"))
        elif kind == "bullet" and (m := re.match(r"^([A-Z][\w /&-]{0,30}):\s+(.+)$", val)):
            flush()
            p = f.para(f"{m[1]}: ", style="List Bullet"); f.inline(p, m[2], m[1])
        else:
            chunk.append(b)
    flush()


def write_doc(f, tpl_md, ex_md, base, product, skip_until=None, stop_at=None, expand_hld=False):
    tsecs = sections(parse(tpl_md))
    esecs = sections(parse(ex_md))
    if skip_until:  # drop the template's own preamble (purpose/process notes)
        start = next(i for i, s in enumerate(tsecs) if s[0] and s[0].startswith(skip_until))
        tsecs = tsecs[start + 1:]
    if stop_at:
        tsecs = tsecs[:next(i for i, s in enumerate(tsecs) if s[0] and s[0].startswith(stop_at))]
    m = match(tsecs, esecs)
    # example sections with no template counterpart fold into the previous matched one
    extra, last = {}, None
    inv = {v: k for k, v in m.items()}
    for ei, (eh, el, eb) in enumerate(esecs):
        if ei in inv: last = inv[ei]
        elif eh and last is not None: extra.setdefault(last, []).extend([("h", el, eh)] + eb)
    for ti, (th, tl, tb) in enumerate(tsecs):
        if th is None: continue
        f.heading(th, {1: 0, 2: 1, 3: 2, 4: 3}[tl] if tl != 1 else 0, product)
        ebody = esecs[m[ti]][2] + extra.get(ti, []) if ti in m else None
        if expand_hld: ebody = builds_on(f, th, ebody)
        render_body(f, tb, ebody, re.sub(r"[*#]", "", th), base)

# System Description sections that expand a High-Level Design section (forms only; the
# markdown templates keep both). Key = section number, value = the HLD heading it builds on.
BUILDS_ON = {"1": "What It Is", "3": "Block Diagram", "4": "Subsystems", "5": "Key Interfaces",
             "9": "Constraints", "10": "Open Calls"}


def builds_on(f, th, ebody):
    """Point an overlapping section at the High-Level Design instead of repeating it."""
    hld = BUILDS_ON.get(key(th))
    if not hld: return ebody
    if hld == "Block Diagram":
        f.para("↳ Builds on your High-Level Design block diagram. Paste an updated diagram here only if it "
               "has changed; otherwise describe the architecture below.", style="Guidance")
        return [b for b in ebody or [] if b[0] not in ("img", "code")]
    f.para(f"↳ Builds on your High-Level Design ({hld}). Start from what you wrote there and add the "
           "detail here.", style="Guidance")
    return ebody

# ---------------------------------------------------------------- checklist

def checklist(f, tpl_md, ex_md):
    esecs = {}
    cur = None
    for kind, _, val in parse(ex_md):
        if kind == "h": cur = norm(val.split("—")[-1])
        elif kind == "table" and cur and val and val[0][:2] == ["#", "Item"]:
            esecs.setdefault(cur, []).extend(val[1:])
    cur, items = None, []

    def emit():
        if not items: return
        ex = esecs.get(norm(cur.split("—")[-1]), [])
        rows = [["", "Item", "Status", "Notes / evidence"]]
        t = f.doc.add_table(rows=1, cols=4); t.style = "Table Grid"
        for c, h in zip(t.rows[0].cells, rows[0]): runs(c.paragraphs[0], h, bold=True)
        for it in items:
            plain = it.replace("**[HW↔SW]**", "").strip()
            best = max(ex, key=lambda r: difflib.SequenceMatcher(None, plain.lower(), r[1].lower()).ratio(), default=None)
            ok = best and difflib.SequenceMatcher(None, plain.lower(), best[1].lower()).ratio() > 0.3
            c = t.add_row().cells
            el, content = sdt("Done", checkbox=True)
            r = OxmlElement("w:r"); rpr = OxmlElement("w:rPr"); fn = OxmlElement("w:rFonts")
            for a in ("w:ascii", "w:hAnsi", "w:eastAsia"): fn.set(qn(a), "MS Gothic")
            rpr.append(fn); r.append(rpr); tx = OxmlElement("w:t"); tx.text = "☐"; r.append(tx); content.append(r)
            c[0].paragraphs[0]._p.append(el)
            runs(c[1].paragraphs[0], it)
            f.inline(c[2].paragraphs[0], best[2] if ok else "Status", "Status", dropdown=["PASS", "FAIL", "N/A"])
            f.inline(c[3].paragraphs[0], best[3] if ok and len(best) > 3 else "Where the system description covers this", "Notes")
        for row in t.rows:
            for c, w in zip(row.cells, (Cm(0.9), Cm(7.6), Cm(1.8), Cm(6.5))): c.width = w
        f.doc.add_paragraph()
        items.clear()

    blocks = parse(tpl_md)
    first = True
    for kind, lvl, val in blocks:
        if kind == "h" and lvl == 1:
            f.heading(val, 0)
        elif kind == "h":
            emit(); cur = val; f.heading(val, lvl - 1)
        elif kind == "check":
            items.append(val)
        elif kind == "p" and cur is None and first:
            f.para(val, style="Guidance")
    emit()

# ---------------------------------------------------------------- main

# What each document is, who writes and reads it, and when. Kept in step with the README.
ABOUT = {
    "high-level-design":
        "The High-Level Design is a one-to-two-page overview of the product: what it is, the main blocks "
        "and how they connect, the constraints that rule options out, and the three hardest problems. The "
        "product manager writes it with one hardware lead and one software lead, in half a day to a day, "
        "after the idea has been explored and before anyone picks parts. Executives, partners, investors "
        "and new team members read it to understand in five minutes what is being built. It is also the "
        "first gate: if it does not hold together, do not start the System Description. "
        "Next: the System Description expands each of its sections.",
    "system-description":
        "The System Description is the full engineering spec: real part numbers, interfaces, power budget, "
        "firmware architecture, connectivity, key decisions, and constraints such as BOM cost, lead times, "
        "minimum order quantities, certification and manufacturing. It starts from the High-Level Design; "
        "sections marked \"Builds on your High-Level Design\" pick up where that document stopped. The "
        "technical lead or systems engineer writes it with the electrical, mechanical, firmware and app "
        "leads, over several days, once the High-Level Design is agreed and before schematics or tooling. "
        "Engineers design from it; suppliers and contract manufacturers quote from it. Keep it current "
        "through the first engineering prototype build (EVT). Next: check it against the Gate Checklist.",
    "gate-checklist":
        "The Gate Checklist has 100 items that test whether the System Description is complete enough to "
        "start detailed design. Mark each item PASS, FAIL or N/A and note where the System Description "
        "covers it. Items tagged [HW↔SW] target the line between hardware and software, where most "
        "integration surprises start. Run it when you believe the System Description is ready, preferably "
        "with someone who did not write it: a reviewer, the engineering manager, or the hardware and "
        "software leads together in a design review. Every FAIL is either a gap to fill in the System "
        "Description or a scope decision to write down. Revise and re-run until the FAILs left are ones "
        "you accept.",
}


def cover(f, product, names):
    f.para().paragraph_format.space_before = Pt(120)
    f.para("ELECTRUM", bold=True, color=ACCENT)
    t = f.doc.add_paragraph(style="Title"); f.inline(t, product, "Product name")
    f.para("Hardware–software partitioning workbook" if len(names) > 1 else names[0], color=GUIDE, size=Pt(14))
    for label, val in (("Date", "YYYY-MM-DD"), ("Author", "Name"), ("Status", "Draft")):
        p = f.para(f"**{label}:** "); f.inline(p, val, label)
    f.para()
    f.para("Grey text is a worked example (Bubbler, an automated soap-bubble maker). Click it and type or paste: "
           "the example disappears. Tick checklist boxes with a click.", style="Guidance")
    if len(names) > 1:
        f.para("This workbook holds all three Electrum documents in one file, in the order you fill them in: "
               "the High-Level Design, then the System Description, then the Gate Checklist. It is not a "
               "separate step, and the other documents are not generated from it: the three single files "
               "contain the same forms. Use the workbook when one person or team takes the product through "
               "all three. Use the single files when different people own each document or you share them "
               "at different times.")
        f.para("**Contents**")
        for i, n in enumerate(names, 1): f.para(f"{i}. {n}")
    else:
        f.para(f"One of three Electrum documents: the High-Level Design comes first, the System Description "
               f"expands it, and the Gate Checklist tests the System Description. The workbook "
               f"(electrum-workbook.docx) holds all three in one file.")


def build(out, example):
    base = Path(example)
    ex = lambda n: (base / n).read_text()
    product = re.match(r"#\s*System Description:\s*(.+)", ex("system_description.md"))[1].strip()
    parts = {
        "high-level-design": ("High-Level System Design", lambda f: write_doc(
            f, (TPL / "hw_sw_high_level.md").read_text(), ex("high_level_design.md"), base, product,
            skip_until="Template", stop_at="What This Is NOT")),
        "system-description": ("System Description", lambda f: write_doc(
            f, (TPL / "system_description_template.md").read_text(), ex("system_description.md"), base, product,
            expand_hld=True)),
        "gate-checklist": ("Gate Checklist", lambda f: checklist(
            f, (TPL / "checklist.md").read_text(), ex("gate_checklist.md"))),
    }
    out = Path(out); out.mkdir(parents=True, exist_ok=True)
    for name, chosen in [("workbook", list(parts))] + [(k, [k]) for k in parts]:
        f = Form()
        cover(f, product, [parts[k][0] for k in chosen])
        for k in chosen:
            f.new_page, f.about = True, ABOUT[k]
            parts[k][1](f)
        path = out / f"electrum-{name}.docx"
        f.doc.save(path)
        print(path)


if __name__ == "__main__":
    build(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else ROOT / "examples/bubbler-automated-soap-bubble-maker")
